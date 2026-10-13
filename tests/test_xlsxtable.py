"""엑셀 → norm CSV 변환기 — **형식만 바꿨는가. 값을 만들지 않았는가.**

★ 2026-10-13 (DECISIONS §515). 기존 대장 80종이 기계로 읽는 형식은 csv · zip ·
  json · txt 넷뿐이고 xlsx 는 0번 나온다. 정보공개청구 회신 둘이 xlsx 로 왔고
  `kind: csv_table` 로 적으면 그 불변식이 깨진다 — 집의 답은 `raw_only` +
  `norm_convert` 다(`gjfire_vehicle_card` 가 PDF 로 먼저 밟았다).

여기가 드는 판별식 넷. 전부 **실측에서 온 카나리아**다.
  ① 꼬리의 빈 행은 떼고 **사이의 빈 행은 남긴다** — 화재 표의 빈 행 셋은
     연도 구분자고, 지우면 파일이 적은 것을 바꾼다.
  ② 병합 칸을 **채우지 않는다** — 일방통행 표의 머리글이 3행 병합인데
     채우면 변환기가 값을 만든다. 범위는 `_prep.json` 에 적는다.
  ③ 숫자는 **저장값**이다 — `#,##0` 이 `3000` 을 `3,000` 으로 보여주는 것은
     화면이다(천변우로 설월교~발산교 D28 이 실제로 그렇다).
  ④ `_x000D_` 는 원문에 없는 글자다 — 엑셀의 CR 벗어남이고 `load_workbook`
     이 안 되돌린다(화재 표 `발화개요` 12건 중 7곳).

★ 실물 레이크가 없으면(CI) 실물 검사는 건너뛴다. 판별식은 합성 시트가 지킨다.

밖    **옮긴 표의 내용이 맞는가는 안 본다** — 33구간이 정말 33인가 · 도로명이
      실제 동구의 것인가는 사람이 읽고 대장 `note` 가 적는다. 대장 항목의
      필드는 `ledger_check` 소관이고 변환기 등록표는 `prep` 소관이다.
      여기가 드는 것은 **형식을 바꾸다가 값이 바뀌었는가** 하나다.
"""
from __future__ import annotations

import datetime as dt

import pytest

from firelane import xlsxtable as xt


def _book(rows, *, fmts=None, merged=(), sheets=("Sheet1",)):
    """합성 통합문서를 메모리에 굽는다. 실물 xlsx 로 왕복해 읽는 쪽을 민다."""
    import io

    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active.title = sheets[0]
    for extra in sheets[1:]:
        wb.create_sheet(extra)
    ws = wb[sheets[0]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
            if fmts and (r, c) in fmts:
                ws.cell(r, c).number_format = fmts[(r, c)]
    for m in merged:
        ws.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf, data_only=True)[sheets[0]]


# ── ① 빈 행 ─────────────────────────────────────────────────────
def test_trailing_blanks_go_and_inner_blanks_stay():
    ws = _book([["a", "b"], [None, None], ["c", "d"], [None, None], [None, None]])
    assert xt.sheet_rows(ws) == [["a", "b"], ["", ""], ["c", "d"]]


def test_trailing_blank_columns_go():
    ws = _book([["a", None, None], ["c", None, None]])
    assert xt.sheet_rows(ws) == [["a"], ["c"]]


def test_a_sheet_of_nothing_is_no_rows():
    assert xt.sheet_rows(_book([[None]])) == []


# ── ② 병합 ──────────────────────────────────────────────────────
def test_merged_cells_are_not_filled_in():
    """일방통행 표의 머리글 꼴 — `연번` 이 A1:A3 병합이다."""
    ws = _book([["연번", "구간"], [None, "길이"], [None, "(m)"]],
               merged=("A1:A3",))
    assert xt.sheet_rows(ws) == [["연번", "구간"], ["", "길이"], ["", "(m)"]]


# ── ③ 값 ────────────────────────────────────────────────────────
@pytest.mark.parametrize(("value", "fmt", "want"), [
    ("대의동77", "General", "대의동77"),      # 원문의 공백 없음을 그대로
    ("관할 ", "General", "관할 "),            # 꼬리 공백도 그대로
    (3000, "#,##0", "3000"),                  # 화면은 `3,000` — 저장값을 쓴다
    (1.5, "General", "1.5"),
    (0, "General", "0"),
    (None, "General", ""),
    (True, "General", "TRUE"),
])
def test_cell_text_moves_the_letters_only(value, fmt, want):
    assert xt.cell_text(value, fmt) == want


@pytest.mark.parametrize(("value", "fmt"), [
    (0.15, "0%"),                                   # 화면 `15%` ≠ 저장값 0.15
    # ★ `tzinfo` 는 엄격 린트(DTZ001)가 요구한다. 엑셀 칸은 시간대가 없지만
    #   이 단이 보는 것은 **갈래가 datetime 인가** 하나다 — 시각은 안 본다.
    (dt.datetime(2022, 2, 12, 13, 3, tzinfo=dt.UTC), "yyyy-mm-dd"),  # 저장값은 일련수
    (dt.date(2022, 2, 12), "yyyy-mm-dd"),
    ("=SUM(A1:A9)", "General"),                     # 캐시가 없으면 조용히 빈다
    (object(), "General"),
])
def test_cell_text_cries_instead_of_choosing(value, fmt):
    with pytest.raises(xt.SheetError):
        xt.cell_text(value, fmt)


def test_the_deceiving_format_canary_is_alive():
    """래칫 — `DECEIVING_FMT` 를 비우면 이 시험이 운다."""
    assert xt.DECEIVING_FMT, "속이는 서식 목록이 비었다 — 백분율이 조용히 지나간다"


# ── ④ 원문에 없는 글자 ──────────────────────────────────────────
def test_excel_escape_is_undone():
    """화재 표 `발화개요` 의 실제 꼴(2026-10-13 · 7곳)."""
    got = xt.cell_text("화재발생_x000D_\n- 지붕 샌드위치판넬  66㎡ 소손_x000D_\n")
    assert "_x000D_" not in got
    assert got == "화재발생\n- 지붕 샌드위치판넬  66㎡ 소손\n"


def test_an_escaped_underscore_stays_a_letter():
    """`_x005F_` 가 `_` 의 벗어남이다 — 사람이 적은 `_x000D_` 는 글자로 남는다."""
    assert xt.cell_text("_x005F_x000D_") == "_x000D_"


# ── 시트가 둘이면 고르지 않는다 ─────────────────────────────────
def test_two_sheets_cry(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.create_sheet("둘째")
    wb.active["A1"] = "값"
    p = tmp_path / "둘.xlsx"
    wb.save(p)
    with pytest.raises(xt.SheetError, match="시트가 2개"):
        xt.read(p)


# ── 산출물 ──────────────────────────────────────────────────────
def test_convert_writes_utf8_lf_and_tells_what_it_moved(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "머리", "값"
    ws["A2"], ws["B2"] = "가", 7
    ws.merge_cells("A1:B1")
    src, dst = tmp_path / "표.xlsx", tmp_path / "표.csv"
    wb.save(src)
    meta = xt.convert(src, dst)
    raw = dst.read_bytes()
    assert b"\r" not in raw
    assert raw.decode("utf-8") == "머리,\n가,7\n"
    assert meta["converter"] == "xlsxtable"
    assert meta["converter_version"] == xt.VERSION
    assert meta["dst_newline"] == "lf"
    assert meta["merged"] == ["A1:B1"], "병합 범위를 안 적으면 지식을 버린다"
    assert meta["sheet_rows"] == 2


def test_the_converter_is_in_the_registry():
    """대장이 `norm_convert: xlsxtable` 을 적어도 표 밖이면 prep 이 터진다."""
    from firelane.prep import CONVERTERS
    assert CONVERTERS["xlsxtable"] == (".xlsx", "firelane.xlsxtable")
    assert hasattr(xt, "convert") and hasattr(xt, "VERSION")

# ── 갈래가 여는 확장자 (DECISIONS §515) ─────────────────────────
# ★ 이 셋은 `tests/test_guards.py` 의 kind 등록부 검사와 **같은 족**이지만
#   그 파일이 `sizetable` 예외 상한(2,911줄)에 닿아 있다. 예외를 올리지 않고
#   여기 둔다 — §511-2 가 `dms.py` 를 쪼갠 것과 같은 규율이다.
def test_only_the_unread_kind_declines_to_name_its_extensions():
    from firelane.kinds import KINDS
    assert {k for k, v in KINDS.items() if v.exts is None} == {"raw_only"}, (
        "`exts` 가 비는 갈래는 `raw_only` 하나다 — 읽지 않는 것만 확장자를 "
        "안 고른다. 새 갈래를 넣었으면 그 갈래가 여는 확장자를 적어라.")


def test_the_machine_reads_exactly_four_formats():
    """실측 2026-10-13 — csv · zip · json · txt. **다섯째는 갈래 추가다.**"""
    from firelane.kinds import READABLE_EXT
    assert READABLE_EXT == {"csv", "zip", "json", "txt"}, (
        f"기계로 읽는 확장자가 바뀌었다 — {sorted(READABLE_EXT)}. "
        "다섯째가 생기는 날은 **갈래를 추가하는 날**이고 `kinds.py` 머리말의 "
        "네 단계를 밟는다. 서식(xlsx·pdf·hwpx…)을 읽고 싶으면 갈래가 아니라 "
        "**변환기**다 — `raw_only` + `norm_convert` 에 `prep.CONVERTERS` 한 줄.")


def test_a_kind_that_cannot_open_that_extension_is_a_false_declaration():
    """카나리아 — 2026-10-13 의 실제 선언(`csv_table` + `ext: [xlsx]`)이다."""
    from firelane.ledger_check import FAIL, check_entry
    e = {"kind": "csv_table", "ext": ["xlsx"], "stem": "x",
         "scope": "jngj-dongmyeong", "updated": "2026-01-01", "authority": "a",
         "what": "w", "encoding": "utf-8", "schema": {"columns": []},
         "feeds": [], "feeds_why": "참조용 — x", "used_for": "u"}
    bad = [str(i) for i in check_entry("가짜", e) if i.level == FAIL]
    assert any("csv_table" in b and "xlsx" in b for b in bad), (
        f"거짓 선언을 통과시킨다 — {bad}")
    e2 = {**e, "kind": "raw_only", "norm_convert": "xlsxtable"}
    assert not [i for i in check_entry("진짜", e2) if i.level == FAIL], (
        "`raw_only` + 변환기는 **옳은 선언**인데 울린다")
