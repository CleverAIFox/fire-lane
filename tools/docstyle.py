#!/usr/bin/env python3
"""
docstyle.py — 기획서에 **개요 층**이 있는가. 목차·북마크가 만들어지는가.  (DECISIONS §278-3)

    uv run python tools/docstyle.py            검사 (종료코드 = 결함 수)
    uv run python tools/docstyle.py --tree     잡힌 제목을 층째로 본다
    uv run python tools/docstyle.py --write    개요 층을 넣는다 (기획서를 고친다)
    uv run python tools/docstyle.py --selftest ★ 판별식이 살아 있나

── 왜 생겼나 (2026-09-28 실측) ────────────────────────────────
`docs/proposal.docx` 는 **문서 넷 중 유일하게 밖이 읽는 것**인데 —

    본문 522단락 중 `pStyle` 이 붙은 것      0개
    `styles.xml` 의 Heading 1~6 선언         있다 (아무도 안 쓴다)
    제목 흉내                                 굵은 글씨로만
    web/proposal.pdf 63쪽의 북마크            0개

★ **제목 계층이 사람 눈에만 있다.** 기계에는 평단락 522개뿐이다. 그래서
  자동 목차가 안 만들어지고 63쪽 PDF 에 북마크가 하나도 없다 — 읽는 쪽이
  스크롤로만 찾는다.

★ **Heading 스타일은 안 씌운다.** 씌우면 글꼴·크기·색이 워드 기본값으로
  바뀌어 **제출본 모양이 달라진다.** `w:outlineLvl` 만 넣으면 보이는 것은
  그대로이고 개요만 생긴다 — 목차도 북마크도 그것으로 만들어진다.

★ **Part 마다 계층이 다르다.** Part I·III 은 장을 `1.` 로 쓰고 그 아래를
  `□` 로 쓰는데, **Part II 는 장을 `□` 로 쓰고 `1.` 을 그 아래 항으로 쓴다.**
  꼴 하나로 층을 정하면 Part II 의 부모·자식이 뒤집힌다. 그래서 지도를
  Part 별로 나눠 선언한다 — 짐작하지 않는다.

IN    docs/proposal.docx
OUT   결함 목록 (비면 통과) · `--write` 면 개요 층을 넣은 docx
PARAM HIER · MAX_LEN
밖    **말투 · 어휘는 안 본다** — `tools/tonecheck.py` 소관이다.
      **내용이 산출물과 맞는지도 안 본다** — `tools/docx_check.py` 가 든다.
      **PDF 를 굽지 않는다** — `tools/proposal_pdf.py` 가 굽고, 개요가 있으면
      북마크는 그쪽이 자동으로 만든다. 여기는 **개요가 있는가**만 본다.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "proposal.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

#: 편(篇) 제목. 무조건 0층이다.
PART_RX = r"^Part\s+([IVX]+)\.\s"

#: 편별 제목 꼴 — **바깥에서 안으로**. 차례가 곧 개요 층이다 (편이 0층).
#: ★ Part II 의 차례가 뒤집혀 있는 것은 오타가 아니라 문서의 실제 구조다.
HIER: dict[str, tuple[str, ...]] = {
    "I": (r"^\d+\.\s", r"^□\s", r"^[a-d]\.\s", r"^\[.+\]$"),
    "II": (r"^□\s", r"^\d+\.\s"),
    #: Part III 의 3층은 두 꼴이다 — `[소제목]` 과 `[ 표이름 ] 설명`.
    "III": (r"^\d+\.\s", r"^□\s", r"^\[.+\]$|^\[\s*[\w/ ]+\]\s+\S"),
}
#: 이보다 긴 줄은 제목이 아니라 문장이다. 굵게 쓴 강조 문장이 있다.
MAX_LEN = 70


def _bold_only(p) -> bool:
    rs = [r for r in p.runs if r.text.strip()]
    return bool(rs) and all(r.bold for r in rs)


def part_of(text: str) -> str | None:
    m = re.match(PART_RX, text.strip())
    return m.group(1) if m else None


def level_of(text: str, part: str | None) -> int | None:
    """이 줄의 개요 층. 제목 꼴이 아니면 None. 편 밖이면 편 제목만 잡는다."""
    s = text.strip()
    if not s or len(s) > MAX_LEN:
        return None
    if part_of(s):
        return 0
    if part is None:
        return None
    for depth, rx in enumerate(HIER.get(part, ()), start=1):
        if re.match(rx, s):
            return depth
    return None


def survey(doc) -> tuple[list, list, list]:
    """(개요가 빠진 제목, 이미 있는 제목, 꼴 밖인데 제목처럼 보이는 줄)."""
    miss, have, odd = [], [], []
    part = None
    for i, p in enumerate(doc.paragraphs):
        s = p.text.strip()
        if not s:
            continue
        lvl = level_of(s, part)
        if lvl is None:
            if _bold_only(p) and len(s) <= MAX_LEN:
                odd.append((i, part, s))
            continue
        part = part_of(s) or part
        if not _bold_only(p):
            continue
        pr = p._p.find(f"{W}pPr")
        cur = pr.find(f"{W}outlineLvl") if pr is not None else None
        (have if cur is not None else miss).append((i, lvl, s))
    return miss, have, odd


def _undeclared(doc) -> list[str]:
    """문서에 있는데 `HIER` 에 없는 편. 짐작으로 층을 주지 않는다."""
    seen = {q for p in doc.paragraphs if (q := part_of(p.text))}
    return sorted(seen - set(HIER))


def _dead(doc) -> list[str]:
    """`HIER` 에 선언했는데 실물에서 한 번도 안 걸린 꼴. 죽은 선언이다."""
    live: set[tuple[str, str]] = set()
    part = None
    for p in doc.paragraphs:
        s = p.text.strip()
        if not s or len(s) > MAX_LEN:
            continue
        if q := part_of(s):
            part = q
            continue
        if part is None or not _bold_only(p):
            continue
        for rx in HIER.get(part, ()):
            if re.match(rx, s):
                live.add((part, rx))
                break
    return [f"Part {p} 의 `{rx}`" for p, rxs in HIER.items() for rx in rxs
            if (p, rx) not in live]


def _set_level(p, lvl: int) -> None:
    """`w:outlineLvl` 만 넣는다. **글꼴 · 크기 · 색은 안 건드린다.**"""
    from docx.oxml.ns import qn  # noqa: PLC0415
    from docx.oxml.shared import OxmlElement  # noqa: PLC0415
    pr = p._p.get_or_add_pPr()
    if (old := pr.find(f"{W}outlineLvl")) is not None:
        pr.remove(old)
    e = OxmlElement("w:outlineLvl")
    e.set(qn("w:val"), str(lvl))
    pr.append(e)


def _open():
    import docx  # noqa: PLC0415
    return docx.Document(str(DOCX))


def check() -> int:
    if not DOCX.is_file():
        print(f"✗ {DOCX.relative_to(ROOT)} 가 없다 — 못 잰 것을 통과로 세지 않는다")
        return 1
    d = _open()
    miss, have, odd = survey(d)
    bad: list[str] = []
    for q in _undeclared(d):
        bad.append(f"Part {q} 의 계층이 `HIER` 에 없다 — 편마다 장 꼴이 다르다. 읽고 선언하라")
    for s in _dead(d):
        bad.append(f"{s} 가 실물에서 한 번도 안 걸린다 — 죽은 선언이다. 지워라")
    if not (miss or have):
        bad.append("제목이 하나도 안 잡혔다 — `HIER` 가 실물과 갈렸다. 판별식을 의심하라")
    if miss:
        bad.append(f"개요 층이 없는 제목 {len(miss)}/{len(miss) + len(have)}")

    if not bad:
        print(f"✓ 개요 층 — 제목 {len(have)}개 전부에 `outlineLvl` 이 있다 "
              f"(목차 · PDF 북마크가 만들어진다)")
        if odd:
            print(f"  참고 — 꼴 밖인데 굵은 짧은 줄 {len(odd)}개는 제목으로 안 본다 (`--tree`)")
        return 0

    print("✗ " + "\n✗ ".join(bad))
    for _i, lvl, s in miss[:10]:
        print(f"      L{lvl}  {s[:64]}")
    if len(miss) > 10:
        print(f"      … {len(miss) - 10}개 더 (`--tree`)")
    if miss:
        print("\n  제목 계층이 **사람 눈에만** 있다 — 자동 목차도 PDF 북마크도 안 만들어진다.")
        print("  넣는 법:  uv run python tools/docstyle.py --write")
        print("  ★ 글꼴 · 크기 · 색은 안 바뀐다. `w:outlineLvl` 만 넣는다.")
    return len(bad) + len(miss)


def write() -> int:
    d = _open()
    if q := _undeclared(d):
        print(f"✗ Part {', '.join(q)} 의 계층이 선언에 없다 — 짐작으로 층을 주지 않는다")
        return 1
    miss, have, _odd = survey(d)
    if not miss:
        print(f"  넣을 것이 없다 — 제목 {len(have)}개 전부 이미 있다")
        return 0
    for i, lvl, _s in miss:
        _set_level(d.paragraphs[i], lvl)
    d.save(str(DOCX))
    print(f"✓ 개요 층 {len(miss)}개 넣었다 (이미 있던 것 {len(have)})")
    print("  ★ 글꼴 · 크기 · 색은 안 건드렸다. 다음 `proposal_pdf` 굽기부터 북마크가 생긴다.")
    return 0


def tree() -> int:
    miss, have, odd = survey(_open())
    for _i, lvl, s in sorted(miss + have):
        print(f"  {'   ' * lvl}{'·' if lvl else '■'} {s[:76]}")
    if odd:
        print(f"\n  꼴 밖 — 제목으로 안 본다 ({len(odd)}개)")
        for _i, p, s in odd:
            print(f"      [{p or '편 밖'}] {s[:70]}")
    return 0


def selftest() -> int:
    """★ 편마다 꼴이 제 층으로 가는가. 편이 뒤집힌 것을 실제로 구분하는가."""
    bad = []
    if level_of("Part II. 요구사항 분석서", None) != 0:
        bad.append("편 제목을 못 잡는다")
    if level_of("1. 프로젝트 개요", None) is not None:
        bad.append("편 밖의 줄에 층을 준다 — 어느 편인지 모르면 모르는 것이다")
    for part, rxs in HIER.items():
        for depth, rx in enumerate(rxs, start=1):
            if len(rxs) != len(set(rxs)):
                bad.append(f"Part {part} 에 같은 꼴이 두 번 있다")
            _ = rx, depth
    # ★ 같은 꼴이 편마다 다른 층으로 가야 한다 — 이게 이 도구의 존재 이유다.
    if level_of("1. 데이터 수집", "I") == level_of("1. 데이터 수집", "II"):
        bad.append("Part I 과 II 에서 `1.` 이 같은 층으로 간다 — 편별 지도가 죽었다")
    if level_of("□ 어떤 절", "I") == level_of("□ 어떤 절", "II"):
        bad.append("Part I 과 II 에서 `□` 가 같은 층으로 간다 — 편별 지도가 죽었다")
    if level_of("이것은 제목이 아니라 그냥 문장이다. 판정은 꼴과 폭으로 한다.", "I") is not None:
        bad.append("평범한 문장을 제목으로 본다")
    if level_of("1. " + "가" * 90, "I") is not None:
        bad.append(f"{MAX_LEN}자를 넘는 줄을 제목으로 본다 — 굵게 쓴 강조 문장이 있다")
    if level_of("", "I") is not None:
        bad.append("빈 줄을 제목으로 본다")
    if level_of("1. 개요", "없는편") is not None:
        bad.append("선언에 없는 편에 층을 준다 — 모르는 것에 층을 주면 안 된다")
    if not HIER:
        bad.append("`HIER` 가 비었다 — 볼 것이 없으면 통과가 아니다")
    if bad:
        print("★ 자기검사 실패\n  " + "\n  ".join(bad))
        return 1
    print(f"✓ 자기검사 — 편 {len(HIER)}개의 계층이 서로 다르게 잡히고, 문장에는 조용하다")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="기획서 개요 층 검사")
    ap.add_argument("--write", action="store_true", help="개요 층을 넣는다")
    ap.add_argument("--tree", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    if a.tree:
        return tree()
    return write() if a.write else check()


if __name__ == "__main__":
    sys.exit(main())
